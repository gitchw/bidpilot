from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from bidpilot.models import IntelligenceBrief, SourceDiagnostic, TenderQuerySpec, TenderRecord

ACCENT = "2F6BFF"
INK = "172033"
MUTED = "657085"
LIGHT = "EAF0FF"


def safe_filename(raw_query: str, timestamp: datetime) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", raw_query)
    cleaned = re.sub(r"\s+", "", cleaned).strip("._")
    if len(cleaned) > 90:
        cleaned = cleaned[:90]
    return f"{cleaned}_{timestamp.strftime('%Y%m%d%H%M')}.docx"


def _set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    tc_pr.append(shading)


def _set_cell_text(cell, text: str, *, bold: bool = False, color: str = INK) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.28

    for style_name, size, color in (
        ("Title", 28, INK),
        ("Heading 1", 18, INK),
        ("Heading 2", 13, ACCENT),
        ("Heading 3", 11, INK),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("标擎 BidPilot · 证据优先的招投标情报")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def generate_report(
    spec: TenderQuerySpec,
    records: list[TenderRecord],
    diagnostics: list[SourceDiagnostic],
    output_dir: Path,
    *,
    generated_at: datetime | None = None,
    incremental: bool = False,
    intelligence_brief: IntelligenceBrief | None = None,
) -> Path:
    generated_at = generated_at or datetime.now(ZoneInfo(spec.schedule.timezone))
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / safe_filename(spec.raw_query, generated_at)

    document = Document()
    _configure_document(document)

    eyebrow = document.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = eyebrow.add_run("BIDPILOT / TENDER INTELLIGENCE")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(ACCENT)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("招投标情报简报")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(spec.raw_query)
    subtitle_run.font.size = Pt(12)
    subtitle_run.font.color.rgb = RGBColor.from_string(MUTED)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        f"生成时间 {generated_at.strftime('%Y-%m-%d %H:%M')}  ·  "
        f"{'增量报告' if incremental else '全量报告'}  ·  共 {len(records)} 条"
    )
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor.from_string(MUTED)

    document.add_paragraph()
    metric_table = document.add_table(rows=2, cols=4)
    metric_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    metric_table.autofit = True
    metrics = [
        ("主题", spec.topic),
        ("地域", spec.region or "全国"),
        ("时间范围", f"{spec.start_date} 至 {spec.end_date}"),
        ("结果", f"{len(records)} 条"),
    ]
    for index, (label, value) in enumerate(metrics):
        _set_cell_text(metric_table.cell(0, index), label, bold=True, color=ACCENT)
        _set_cell_shading(metric_table.cell(0, index), LIGHT)
        _set_cell_text(metric_table.cell(1, index), value, bold=True)

    document.add_heading("执行摘要", level=1)
    if records:
        high = sum(1 for record in records if record.opportunity_score >= 80)
        projects = len({record.lifecycle_id for record in records})
        duplicates = sum(max(0, record.duplicate_count - 1) for record in records)
        paragraph = document.add_paragraph()
        paragraph.add_run(
            f"本次从 {len(diagnostics)} 个来源通道检索并保留 {len(records)} 条相关公告，"
            f"归属 {projects} 个项目生命周期；高优先级机会 {high} 条，跨站重复合并 {duplicates} 条。"
        )
        if incremental:
            paragraph.add_run(" 本报告仅包含该订阅尚未成功投递的新公告或新版本。").bold = True
    else:
        document.add_paragraph(
            "本次在已成功访问的来源与查询口径内未发现匹配结果。请结合下方来源覆盖状态判断是否需要扩大时间、地域或关键词。"
        )

    if intelligence_brief:
        document.add_heading("情报副驾驶", level=1)
        status_labels = {
            "applied": "证据约束 AI",
            "cached": "证据约束 AI（缓存复用）",
            "disabled": "确定性分析（AI 已关闭）",
            "not_configured": "确定性分析（模型未配置）",
            "invalid_response": "确定性分析（模型输出未通过校验）",
            "unavailable": "确定性分析（模型不可用）",
            "empty": "零结果诊断",
        }
        brief_meta = document.add_paragraph()
        meta_label = brief_meta.add_run(
            status_labels.get(intelligence_brief.status, intelligence_brief.status)
        )
        meta_label.bold = True
        meta_label.font.color.rgb = RGBColor.from_string(ACCENT)
        brief_meta.add_run(f"  ·  {intelligence_brief.summary}")
        document.add_paragraph(intelligence_brief.overview)

        if intelligence_brief.priorities:
            document.add_heading("优先机会", level=2)
            priority_table = document.add_table(rows=1, cols=4)
            priority_table.style = "Table Grid"
            priority_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for index, text in enumerate(("证据与机会", "阶段 / 分数", "优先原因", "建议动作")):
                _set_cell_text(priority_table.cell(0, index), text, bold=True, color="FFFFFF")
                _set_cell_shading(priority_table.cell(0, index), ACCENT)
            for priority in intelligence_brief.priorities:
                cells = priority_table.add_row().cells
                _set_cell_text(cells[0], f"[{priority.evidence_id}] {priority.title}")
                _add_hyperlink(cells[0].paragraphs[0], "打开原文", priority.source_url)
                _set_cell_text(
                    cells[1],
                    f"{priority.event_type}\n{priority.opportunity_score:.1f}",
                    bold=True,
                )
                _set_cell_text(cells[2], priority.reason)
                _set_cell_text(cells[3], priority.recommended_action)

        if intelligence_brief.buyer_needs:
            document.add_heading("采购需求判断", level=2)
            for claim in intelligence_brief.buyer_needs:
                document.add_paragraph(
                    f"[{', '.join(claim.evidence_ids)}] {claim.text}",
                    style="List Bullet",
                )

        if intelligence_brief.risks:
            document.add_heading("风险提示", level=2)
            risk_labels = {"high": "高", "medium": "中", "low": "低"}
            for risk in intelligence_brief.risks:
                document.add_paragraph(
                    f"[{risk_labels[risk.level]}] [{', '.join(risk.evidence_ids)}] {risk.text}",
                    style="List Bullet",
                )

        if intelligence_brief.actions:
            document.add_heading("下一步行动", level=2)
            for action in intelligence_brief.actions:
                document.add_paragraph(
                    f"[{action.priority}] [{', '.join(action.evidence_ids)}] {action.text}",
                    style="List Bullet",
                )

    document.add_heading("查询口径", level=1)
    scope = document.add_table(rows=0, cols=2)
    scope.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in (
        ("原始问题", spec.raw_query),
        ("主题与扩展词", "、".join(spec.keywords)),
        ("地域", spec.region or "全国"),
        ("时间", f"{spec.start_date} 至 {spec.end_date}"),
        ("执行计划", spec.schedule.expression),
        ("解析器", spec.parser_version),
    ):
        cells = scope.add_row().cells
        _set_cell_text(cells[0], label, bold=True, color=ACCENT)
        _set_cell_shading(cells[0], LIGHT)
        _set_cell_text(cells[1], value)

    document.add_heading("来源覆盖与质量", level=1)
    coverage = document.add_table(rows=1, cols=6)
    coverage.style = "Table Grid"
    coverage.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, text in enumerate(("来源", "状态", "扫描", "候选", "保留", "淘汰与说明")):
        _set_cell_text(coverage.cell(0, index), text, bold=True, color="FFFFFF")
        _set_cell_shading(coverage.cell(0, index), ACCENT)
    for diagnostic in diagnostics:
        cells = coverage.add_row().cells
        reason_labels = {
            "outside_time": "时间外",
            "region_mismatch": "地域不符",
            "buyer_mismatch": "采购单位不符",
            "event_type_mismatch": "类型不符",
            "excluded_keyword": "命中排除词",
            "keyword_mismatch": "主题未命中",
            "low_relevance": "相关度不足",
        }
        rejection = "；".join(
            f"{reason_labels.get(reason, reason)} {count}"
            for reason, count in diagnostic.rejection_reasons.items()
        )
        values = (
            diagnostic.source,
            diagnostic.status.value,
            str(diagnostic.scanned_count),
            str(diagnostic.fetched_count),
            str(diagnostic.kept_count),
            "；".join(item for item in (rejection, diagnostic.message) if item) or "-",
        )
        for index, value in enumerate(values):
            _set_cell_text(cells[index], value)

    document.add_heading("重点标讯", level=1)
    for index, record in enumerate(records, start=1):
        heading = document.add_heading(level=2)
        heading.add_run(f"{index:02d}  {record.title}")

        info = document.add_table(rows=2, cols=4)
        info.alignment = WD_TABLE_ALIGNMENT.CENTER
        values = [
            ("发布时间", record.published_at.strftime("%Y-%m-%d")),
            ("地域", record.region or "未标注"),
            ("采购人", record.buyer or "原文未明确"),
            ("公告类型", record.event_type.value),
            ("机会分", f"{record.opportunity_score:.1f}"),
            ("相关分", f"{record.relevance_score:.1f}"),
            ("项目编号", record.project_id or "未提取"),
            ("跨站合并", f"{record.duplicate_count} 条来源记录"),
        ]
        for value_index, (label, value) in enumerate(values):
            row, column = divmod(value_index, 4)
            cell = info.cell(row, column)
            _set_cell_text(cell, f"{label}\n{value}", bold=value_index < 4)
            if row == 0:
                _set_cell_shading(cell, "F4F7FC")

        summary = document.add_paragraph()
        label = summary.add_run("核心内容  ")
        label.bold = True
        label.font.color.rgb = RGBColor.from_string(ACCENT)
        summary.add_run(record.summary)

        sources = document.add_paragraph()
        source_label = sources.add_run("来源链接  ")
        source_label.bold = True
        for source_index, url in enumerate(record.source_urls, start=1):
            if source_index > 1:
                sources.add_run("  ·  ")
            _add_hyperlink(sources, f"来源 {source_index}", url)

        attachments = document.add_paragraph()
        attachment_label = attachments.add_run("附件链接  ")
        attachment_label.bold = True
        if record.attachments:
            for attachment_index, attachment in enumerate(record.attachments, start=1):
                if attachment_index > 1:
                    attachments.add_run("  ·  ")
                _add_hyperlink(attachments, attachment.name[:30], attachment.url)
        else:
            attachments.add_run("原文未发现可公开下载附件")

        if record.evidence:
            evidence = document.add_paragraph()
            evidence.paragraph_format.left_indent = Cm(0.6)
            evidence.paragraph_format.right_indent = Cm(0.4)
            evidence_run = evidence.add_run(f"证据片段：{record.evidence[0].text[:260]}")
            evidence_run.italic = True
            evidence_run.font.color.rgb = RGBColor.from_string(MUTED)

    document.add_heading("可信性与使用说明", level=1)
    for text in (
        "本报告只汇总本次成功访问且满足查询口径的来源；来源失败或登录过期已在覆盖表中披露。",
        "摘要由原文证据约束生成，数字、日期、机构名无法回指原文时自动使用抽取式摘要。",
        "跨站重复公告会合并来源；同一项目的意向、招标、更正、中标和合同仍作为不同生命周期事件保留。",
        "投标决策前请打开来源链接核验最新公告、资格条件、截止时间和附件。",
    ):
        document.add_paragraph(text, style="List Bullet")

    document.add_section(WD_SECTION.CONTINUOUS)
    document.save(path)
    return path
