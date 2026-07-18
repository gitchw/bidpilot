from datetime import datetime
from pathlib import Path

from docx import Document

from bidpilot.config import Settings
from bidpilot.db import Database
from bidpilot.models import (
    Attachment,
    EventType,
    EvidenceSpan,
    RawTender,
    SourceDiagnostic,
    SourceStatus,
)
from bidpilot.normalize import deduplicate_records, lifecycle_groups, normalize_item
from bidpilot.report import generate_report
from bidpilot.summarize import EvidenceSummarizer


def raw_tender(source: str, url: str, title: str, body: str, event=EventType.TENDER):
    return RawTender(
        source=source,
        source_url=url,
        title=title,
        published_at=datetime(2026, 7, 10, 9, 30),
        region="安徽",
        buyer="安徽大学",
        body=body,
        event_type=event,
        project_id="AH-2026-001",
        attachments=[Attachment(name="采购文件", url=f"{url}/file.docx")],
        evidence=[EvidenceSpan(text=body, source_url=url)],
    )


async def test_normalization_deduplicates_cross_site_but_keeps_lifecycle(sample_spec):
    settings = Settings()
    summarizer = EvidenceSummarizer(settings)
    body = "项目编号：AH-2026-001。预算金额：1200 万元。采购 20 台 GPU 服务器。"
    first = await normalize_item(
        raw_tender("中国政府采购网", "https://a.example/1", "安徽大学 GPU 服务器采购公告", body),
        sample_spec,
        summarizer,
    )
    mirror = await normalize_item(
        raw_tender("中国招标投标网", "https://b.example/2", "安徽大学GPU服务器采购招标公告", body),
        sample_spec,
        summarizer,
    )
    award = await normalize_item(
        raw_tender(
            "中国政府采购网",
            "https://a.example/3",
            "安徽大学 GPU 服务器采购中标公告",
            body + "中标供应商为某科技公司。",
            EventType.AWARD,
        ),
        sample_spec,
        summarizer,
    )
    records = deduplicate_records([first, mirror, award])
    assert len(records) == 2
    tender = next(record for record in records if record.event_type == EventType.TENDER)
    assert tender.duplicate_count == 2
    assert len(tender.sources) == 2
    assert len(lifecycle_groups(records)) == 1


async def test_irrelevant_region_is_filtered(sample_spec):
    item = raw_tender(
        "测试源",
        "https://example.com/1",
        "北京大学 GPU 服务器采购公告",
        "北京地区服务器采购。",
    )
    item.region = "北京"
    item.buyer = "北京大学"
    result = await normalize_item(item, sample_spec, EvidenceSummarizer(Settings()))
    assert result is None


async def test_exclude_keywords_and_event_types_are_enforced(sample_spec):
    sample_spec.exclude_keywords = ["运维服务"]
    sample_spec.event_types = [EventType.TENDER]
    excluded = raw_tender(
        "测试源",
        "https://example.com/excluded",
        "安徽大学 GPU 服务器运维服务采购公告",
        "本项目采购 GPU 服务器运维服务。",
    )
    award = raw_tender(
        "测试源",
        "https://example.com/award",
        "安徽大学 GPU 服务器采购中标公告",
        "GPU 服务器采购结果。",
        EventType.AWARD,
    )
    accepted = raw_tender(
        "测试源",
        "https://example.com/accepted",
        "安徽大学 GPU 服务器采购公告",
        "本项目采购 20 台 GPU 服务器。",
    )
    summarizer = EvidenceSummarizer(Settings())
    assert await normalize_item(excluded, sample_spec, summarizer) is None
    assert await normalize_item(award, sample_spec, summarizer) is None
    assert await normalize_item(accepted, sample_spec, summarizer) is not None


async def test_report_contains_required_fields_and_filename(sample_spec, tmp_path: Path):
    body = "项目编号：AH-2026-001。预算金额：1200 万元。投标截止时间：2026-07-30。"
    record = await normalize_item(
        raw_tender("中国政府采购网", "https://example.com/notice", "安徽大学服务器采购公告", body),
        sample_spec,
        EvidenceSummarizer(Settings()),
    )
    diagnostics = [
        SourceDiagnostic(
            source="中国政府采购网",
            status=SourceStatus.OK,
            fetched_count=2,
            kept_count=1,
            message="完成",
        )
    ]
    generated_at = datetime(2026, 7, 17, 14, 24)
    path = generate_report(sample_spec, [record], diagnostics, tmp_path, generated_at=generated_at)
    assert path.name.endswith("_202607171424.docx")
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    combined = text + table_text
    for required in ("安徽大学服务器采购公告", "发布时间", "来源链接", "核心内容", "附件链接"):
        assert required in combined


def test_delivery_ledger_is_exactly_once(tmp_path: Path):
    database = Database(tmp_path / "test.db")
    # Use a complete real spec for persistence.
    from bidpilot.intent import IntentParser

    spec = IntentParser().parse("最近1个月安徽服务器招标信息", now=datetime(2026, 7, 17))
    database.create_subscription("sub-1", "每日简报", spec, "local", None)
    keys = [("canonical-1", "version-1"), ("canonical-2", "version-1")]
    assert database.undelivered_keys("sub-1", keys) == set(keys)
    database.mark_delivered("sub-1", keys, "report.docx")
    assert database.undelivered_keys("sub-1", keys) == set()
    assert database.undelivered_keys("sub-1", [*keys, ("canonical-1", "version-2")]) == {
        ("canonical-1", "version-2")
    }


def test_extractive_summary_does_not_leak_masked_member_values():
    item = raw_tender(
        "千里马",
        "https://example.com/masked",
        "安徽大学服务器采购公告",
        "本项目采购高性能服务器。预算金额：**万元。投标截止时间：****年**月**日。",
    )
    result = EvidenceSummarizer(Settings()).extractive(item)
    assert "*" not in result.summary
    assert "**万元" not in result.summary
    assert "公开摘要存在脱敏字段" in result.summary
    assert "需授权后核验" in result.summary
    assert len(result.summary) <= 260
