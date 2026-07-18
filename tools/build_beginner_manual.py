from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BODY_LATIN = "Calibri"
BODY_CJK = "Microsoft YaHei"
MONO_LATIN = "Consolas"
INK = "333333"
MUTED = "667085"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
BRAND = "0F8F8A"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
BORDER = "CAD5E2"
WHITE = "FFFFFF"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


@dataclass(frozen=True)
class ScreenshotSpec:
    filename: str
    caption: str
    alt: str
    width_inches: float = 6.25


SCREENSHOTS = {
    "6. 网页六个区域": ScreenshotSpec(
        "01-home.png",
        "图 1  情报检索首页：输入自然语言、选择交付方式，再解析或执行",
        "标擎 BidPilot 情报检索首页",
    ),
    "9. 网页配置模型：逐字段解释": ScreenshotSpec(
        "02-config.png",
        "图 2  配置中心：模型地址、名称、模式和阈值均可在网页设置",
        "标擎 BidPilot 配置中心的模型配置区域",
    ),
    "11. 长期订阅：为什么关闭网页后还能跑": ScreenshotSpec(
        "03-subscriptions.png",
        "图 3  订阅中心：查看 worker、下次执行、最近运行和管理操作",
        "标擎 BidPilot 订阅中心",
    ),
    "17.8 抓取结果为 0": ScreenshotSpec(
        "04-zero-results.png",
        "图 4  零结果诊断：扫描、候选、排除原因、覆盖缺口和安全建议",
        "标擎 BidPilot 零结果原因诊断页面",
    ),
}


def set_run_font(
    run,
    *,
    latin: str = BODY_LATIN,
    cjk: str = BODY_CJK,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), cjk)
    fonts.set(qn("w:cs"), latin)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, *, latin: str, cjk: str, size: float, color: str) -> None:
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), cjk)
    fonts.set(qn("w:cs"), latin)


def set_outline_level(style, level: int) -> None:
    ppr = style.element.get_or_add_pPr()
    existing = ppr.find(qn("w:outlineLvl"))
    if existing is not None:
        ppr.remove(existing)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level))
    ppr.append(outline)


def set_paragraph_shading(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    existing = ppr.find(qn("w:shd"))
    if existing is not None:
        ppr.remove(existing)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    ppr.append(shading)


def set_paragraph_border(
    paragraph,
    *,
    side: str = "start",
    color: str = BLUE,
    size: int = 12,
    space: int = 6,
) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    border = borders.find(qn(f"w:{side}"))
    if border is None:
        border = OxmlElement(f"w:{side}")
        borders.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)


def set_cell_shading(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    existing = tcpr.find(qn("w:shd"))
    if existing is not None:
        tcpr.remove(existing)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    tcpr.append(shading)


def set_cell_margins(
    cell, *, top: int = 60, start: int = 120, bottom: int = 60, end: int = 120
) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:tblHeader")) is None:
        trpr.append(OxmlElement("w:tblHeader"))


def add_field(paragraph, instruction: str, cached: str = "1") -> None:
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)

    instruction_run = paragraph.add_run()
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = f" {instruction} "
    instruction_run._r.append(code)

    separate_run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(separate)

    cached_run = paragraph.add_run(cached)
    set_run_font(cached_run, size=9, color=MUTED)

    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def configure_styles(doc: Document) -> None:
    styles = doc.styles

    normal = styles["Normal"]
    set_style_font(normal, latin=BODY_LATIN, cjk=BODY_CJK, size=11, color=INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    heading_specs = {
        # Let Word choose the page boundary instead of forcing every short chapter
        # onto a fresh page. ``keep_with_next`` below still keeps each chapter
        # heading attached to its first paragraph or screenshot.
        "Heading 1": (16, BLUE, 18, 10, 0, False),
        "Heading 2": (13, BLUE, 14, 7, 1, False),
        "Heading 3": (12, DARK_BLUE, 10, 5, 2, False),
    }
    for name, (size, color, before, after, outline, page_break) in heading_specs.items():
        style = styles[name]
        set_style_font(style, latin=BODY_LATIN, cjk=BODY_CJK, size=size, color=color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.page_break_before = page_break
        set_outline_level(style, outline)

    code = styles.add_style("CodeBlock", 1)
    set_style_font(code, latin=MONO_LATIN, cjk=BODY_CJK, size=9, color="24364B")
    code.paragraph_format.left_indent = Inches(0.15)
    code.paragraph_format.right_indent = Inches(0.15)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.15

    callout = styles.add_style("ManualCallout", 1)
    set_style_font(callout, latin=BODY_LATIN, cjk=BODY_CJK, size=10.5, color="344054")
    callout.paragraph_format.left_indent = Inches(0.12)
    callout.paragraph_format.right_indent = Inches(0.12)
    callout.paragraph_format.space_before = Pt(5)
    callout.paragraph_format.space_after = Pt(9)
    callout.paragraph_format.line_spacing = 1.2
    callout.paragraph_format.keep_together = True

    caption = styles.add_style("ManualCaption", 1)
    set_style_font(caption, latin=BODY_LATIN, cjk=BODY_CJK, size=9, color=MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_with_next = False
    caption.paragraph_format.keep_together = True

    toc_heading = styles.add_style("ManualTocHeading", 1)
    set_style_font(toc_heading, latin=BODY_LATIN, cjk=BODY_CJK, size=20, color=DARK_BLUE)
    toc_heading.font.bold = True
    toc_heading.paragraph_format.space_before = Pt(0)
    toc_heading.paragraph_format.space_after = Pt(10)
    toc_heading.paragraph_format.keep_with_next = True


def configure_numbering(doc: Document) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
        if node.get(qn("w:abstractNumId")) is not None
    ]
    start_abstract = max(abstract_ids, default=0) + 1

    def add_abstract(abstract_id: int, kind: str) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "hybridMultilevel")
        abstract.append(multi)

        bullets = ("•", "○", "▪")
        for level in range(3):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(level))
            start = OxmlElement("w:start")
            start.set(qn("w:val"), "1")
            lvl.append(start)
            num_fmt = OxmlElement("w:numFmt")
            num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
            lvl.append(num_fmt)
            lvl_text = OxmlElement("w:lvlText")
            if kind == "bullet":
                lvl_text.set(qn("w:val"), bullets[level])
            else:
                lvl_text.set(qn("w:val"), "%1." if level == 0 else f"%{level + 1}.")
            lvl.append(lvl_text)
            lvl_jc = OxmlElement("w:lvlJc")
            lvl_jc.set(qn("w:val"), "left")
            lvl.append(lvl_jc)
            ppr = OxmlElement("w:pPr")
            spacing = OxmlElement("w:spacing")
            spacing.set(qn("w:after"), "80")
            spacing.set(qn("w:line"), "300")
            spacing.set(qn("w:lineRule"), "auto")
            ppr.append(spacing)
            tabs = OxmlElement("w:tabs")
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "num")
            tab.set(qn("w:pos"), str(540 + level * 360))
            tabs.append(tab)
            ppr.append(tabs)
            indent = OxmlElement("w:ind")
            indent.set(qn("w:left"), str(540 + level * 360))
            indent.set(qn("w:hanging"), "270")
            ppr.append(indent)
            lvl.append(ppr)
            if kind == "bullet":
                rpr = OxmlElement("w:rPr")
                fonts = OxmlElement("w:rFonts")
                fonts.set(qn("w:ascii"), BODY_CJK)
                fonts.set(qn("w:hAnsi"), BODY_CJK)
                fonts.set(qn("w:eastAsia"), BODY_CJK)
                rpr.append(fonts)
                lvl.append(rpr)
            abstract.append(lvl)

        first_num = numbering.find(qn("w:num"))
        if first_num is None:
            numbering.append(abstract)
        else:
            numbering.insert(list(numbering).index(first_num), abstract)

    add_abstract(start_abstract, "bullet")
    add_abstract(start_abstract + 1, "number")
    return start_abstract, start_abstract + 1


def next_numbering_id(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
        if node.get(qn("w:numId")) is not None
    ]
    return max(ids, default=0) + 1


def create_numbering_instance(doc: Document, abstract_id: int, start: int = 1) -> int:
    numbering = doc.part.numbering_part.element
    num_id = next_numbering_id(doc)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    if start != 1:
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        start_override = OxmlElement("w:startOverride")
        start_override.set(qn("w:val"), str(start))
        override.append(start_override)
        num.append(override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int, level: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    existing = ppr.find(qn("w:numPr"))
    if existing is not None:
        ppr.remove(existing)
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    numpr.append(ilvl)
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.append(numid)
    ppr.insert(0, numpr)


INLINE_PATTERN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")


def add_rich_text(paragraph, text: str, *, size: float | None = None) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, size=size, color=INK)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(
                run,
                latin=MONO_LATIN,
                cjk=BODY_CJK,
                size=9.5 if size is None else min(size, 9.5),
                color=DARK_BLUE,
            )
            rpr = run._element.get_or_add_rPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:val"), "clear")
            shading.set(qn("w:color"), "auto")
            shading.set(qn("w:fill"), LIGHT_BLUE)
            rpr.append(shading)
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=INK, bold=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size, color=INK)


def add_separator(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    set_paragraph_border(paragraph, side="bottom", color=BORDER, size=4, space=1)


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph(style="CodeBlock")
    set_paragraph_shading(paragraph, LIGHT_GRAY)
    set_paragraph_border(paragraph, side="start", color=BLUE, size=10, space=6)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_run_font(run, latin=MONO_LATIN, cjk=BODY_CJK, size=9, color="24364B")
        if index < len(lines) - 1:
            run.add_break()


def table_column_widths(rows: list[list[str]]) -> list[int]:
    columns = max(len(row) for row in rows)
    weights: list[float] = []
    for column in range(columns):
        longest = max((len(row[column]) if column < len(row) else 0) for row in rows)
        weights.append(float(min(max(longest + 3, 9), 42)))
    total = sum(weights)
    widths = [round(CONTENT_WIDTH_DXA * weight / total) for weight in weights]
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def set_table_geometry(table, widths: list[int]) -> None:
    tbl = table._tbl
    tblpr = tbl.tblPr

    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.insert(0, tblw)
    tblw.set(qn("w:w"), str(sum(widths)))
    tblw.set(qn("w:type"), "dxa")

    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tblind.set(qn("w:type"), "dxa")

    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for side in ("top", "start", "bottom", "end", "insideH", "insideV"):
        node = borders.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.insert(0, tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    normalized = [row + [""] * (columns - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = table_column_widths(normalized)
    set_table_geometry(table, widths)

    for row_index, row_data in enumerate(normalized):
        row = table.rows[row_index]
        if row_index == 0:
            set_repeat_table_header(row)
        for column_index, text in enumerate(row_data):
            cell = row.cells[column_index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            add_rich_text(paragraph, text.strip(), size=9.5)
            for run in paragraph.runs:
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_screenshot(doc: Document, path: Path, spec: ScreenshotSpec) -> None:
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_together = True
    shape = paragraph.add_run().add_picture(str(path), width=Inches(spec.width_inches))
    shape._inline.docPr.set("descr", spec.alt)
    shape._inline.docPr.set("name", spec.alt)
    caption = doc.add_paragraph(spec.caption, style="ManualCaption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), alignment=WD_TAB_ALIGNMENT.RIGHT)
    left = paragraph.add_run("标擎 BidPilot · 零基础操作说明书")
    set_run_font(left, size=9, color=MUTED, bold=True)
    paragraph.add_run("\t")
    right = paragraph.add_run("v0.7.0")
    set_run_font(right, size=9, color=MUTED)
    set_paragraph_border(paragraph, side="bottom", color=BORDER, size=4, space=2)

    first_header = section.first_page_header
    first_header.paragraphs[0].clear()

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_paragraph.paragraph_format.space_before = Pt(2)
    prefix = footer_paragraph.add_run("第 ")
    set_run_font(prefix, size=9, color=MUTED)
    add_field(footer_paragraph, "PAGE")
    middle = footer_paragraph.add_run(" 页 / 共 ")
    set_run_font(middle, size=9, color=MUTED)
    add_field(footer_paragraph, "NUMPAGES")
    suffix = footer_paragraph.add_run(" 页")
    set_run_font(suffix, size=9, color=MUTED)

    section.first_page_footer.paragraphs[0].clear()


def enable_field_updates(doc: Document) -> None:
    settings = doc.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = OxmlElement("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")


def add_cover(doc: Document) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(92)
    kicker.paragraph_format.space_after = Pt(14)
    kicker.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = kicker.add_run("BIDPILOT · OPERATOR HANDBOOK")
    set_run_font(run, size=10, color=BRAND, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("标擎 BidPilot")
    set_run_font(run, size=29, color=DARK_BLUE, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("零基础操作说明书")
    set_run_font(run, size=23, color=BLUE, bold=True)

    tagline = doc.add_paragraph()
    tagline.paragraph_format.space_after = Pt(24)
    run = tagline.add_run("从第一次安装到长期订阅、模型配置、零结果诊断与安全停服")
    set_run_font(run, size=12.5, color=MUTED)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    rows = [
        ("文档版本", "v0.7.0"),
        ("适用系统", "Windows / macOS / Linux"),
        ("默认地址", "http://127.0.0.1:8000"),
        ("更新日期", date.today().isoformat()),
    ]
    widths = [2700, 6660]
    set_table_geometry(table, widths)
    for row, (label, value) in zip(table.rows, rows, strict=True):
        for cell in row.cells:
            set_cell_margins(cell, top=90, bottom=90)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        p_label = row.cells[0].paragraphs[0]
        p_label.paragraph_format.space_after = Pt(0)
        label_run = p_label.add_run(label)
        set_run_font(label_run, size=10, color=DARK_BLUE, bold=True)
        p_value = row.cells[1].paragraphs[0]
        p_value.paragraph_format.space_after = Pt(0)
        value_run = p_value.add_run(value)
        set_run_font(value_run, size=10, color=INK)

    callout = doc.add_paragraph(style="ManualCallout")
    callout.paragraph_format.space_before = Pt(24)
    set_paragraph_shading(callout, LIGHT_GRAY)
    set_paragraph_border(callout, side="start", color=BRAND, size=14, space=7)
    lead = callout.add_run("适合谁：")
    set_run_font(lead, size=10.5, color=DARK_BLUE, bold=True)
    rest = callout.add_run(
        "即使你从未使用过 Python、终端、API 或 Docker，也可以按图按步骤完成安装、"
        "启动、查询、订阅、停止、备份与排错。"
    )
    set_run_font(rest, size=10.5, color="344054")

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(16)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("证据优先 · 不伪造结果 · 覆盖边界透明 · 配置可视化")
    set_run_font(run, size=10, color=BRAND, bold=True)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def add_front_matter(doc: Document) -> None:
    heading = doc.add_paragraph("目录", style="ManualTocHeading")
    heading.paragraph_format.keep_with_next = True
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "目录将在打开文档时自动更新")
    note = doc.add_paragraph(style="ManualCallout")
    set_paragraph_shading(note, LIGHT_GRAY)
    set_paragraph_border(note, side="start", color=BLUE, size=10, space=6)
    add_rich_text(
        note,
        "如果目录页码没有自动刷新：在 Word 中按 `Ctrl+A`，再按 `F9`；Mac 可使用 `Command+A` 后更新域。左侧“导航窗格”也可按标题直接跳转。",
        size=10,
    )

    route_heading = doc.add_paragraph("第一次阅读路线", style="Heading 2")
    route_heading.paragraph_format.page_break_before = False
    routes = [
        ["你的目标", "直接阅读"],
        ["5 分钟启动并查询", "第 0、4、7 章"],
        ["设置模型和推送", "第 8～10 章"],
        ["创建并管理长期任务", "第 11 章"],
        ["看懂空结果", "第 7.6、17.8 节"],
        ["备份、恢复和卸载", "第 16、20 章"],
        ["网页打不开或任务不跑", "第 17 章"],
    ]
    add_markdown_table(doc, routes)


def parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(line: str) -> bool:
    cells = parse_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def add_body_from_markdown(
    doc: Document,
    markdown: str,
    *,
    assets_dir: Path,
    bullet_abstract_id: int,
    number_abstract_id: int,
) -> None:
    lines = markdown.splitlines()
    first_section = next((index for index, line in enumerate(lines) if line.startswith("## ")), 0)
    lines = lines[first_section:]

    current_list_kind: str | None = None
    current_num_id: int | None = None
    last_explicit_number: int | None = None
    index = 0

    def reset_list() -> None:
        nonlocal current_list_kind, current_num_id, last_explicit_number
        current_list_kind = None
        current_num_id = None
        last_explicit_number = None

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            language = stripped[3:].strip()
            block: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            if language and language not in {"text", "bash", "shell", "json"}:
                block.insert(0, f"[{language}]")
            add_code_block(doc, block)
            index += 1
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            reset_list()
            marks, text = heading_match.groups()
            style_name = {2: "Heading 1", 3: "Heading 2", 4: "Heading 3"}[len(marks)]
            paragraph = doc.add_paragraph(style=style_name)
            add_rich_text(paragraph, text)
            screenshot = SCREENSHOTS.get(text)
            if screenshot:
                add_screenshot(doc, assets_dir / screenshot.filename, screenshot)
            index += 1
            continue

        if stripped == "---":
            reset_list()
            add_separator(doc)
            index += 1
            continue

        if (
            stripped.startswith("|")
            and index + 1 < len(lines)
            and is_table_separator(lines[index + 1])
        ):
            reset_list()
            table_rows = [parse_table_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append(parse_table_row(lines[index]))
                index += 1
            add_markdown_table(doc, table_rows)
            continue

        list_match = re.match(r"^(\s*)([-*+]|(\d+)\.)\s+(.+)$", line)
        if list_match:
            indent, marker, explicit, text = list_match.groups()
            kind = "number" if explicit is not None else "bullet"
            explicit_number = int(explicit) if explicit is not None else None
            new_sequence = kind != current_list_kind or current_num_id is None
            if kind == "number" and explicit_number == 1 and last_explicit_number not in {None, 0}:
                new_sequence = True
            if new_sequence:
                abstract_id = number_abstract_id if kind == "number" else bullet_abstract_id
                current_num_id = create_numbering_instance(
                    doc,
                    abstract_id,
                    start=explicit_number or 1,
                )
                current_list_kind = kind
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.25
            apply_numbering(paragraph, current_num_id, min(len(indent) // 2, 2))
            add_rich_text(paragraph, text)
            if explicit_number is not None:
                last_explicit_number = explicit_number
            index += 1
            continue

        reset_list()
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index]
            candidate_stripped = candidate.strip()
            if not candidate_stripped:
                break
            if candidate_stripped.startswith(("## ", "### ", "#### ", "```", "|", "---")):
                break
            if re.match(r"^(\s*)([-*+]|(\d+)\.)\s+", candidate):
                break
            paragraph_lines.append(candidate_stripped)
            index += 1
        text = " ".join(part.removesuffix("  ") for part in paragraph_lines)
        callout_prefixes = (
            "注意：",
            "重要：",
            "原理：",
            "以下命令会",
            "如果你不知道占用者是什么",
        )
        if text.startswith(callout_prefixes):
            paragraph = doc.add_paragraph(style="ManualCallout")
            set_paragraph_shading(paragraph, LIGHT_GRAY)
            set_paragraph_border(paragraph, side="start", color=BRAND, size=10, space=6)
        else:
            paragraph = doc.add_paragraph()
        add_rich_text(paragraph, text)


def build_manual(source: Path, output: Path, assets_dir: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    bullet_abstract_id, number_abstract_id = configure_numbering(doc)
    enable_field_updates(doc)
    doc.core_properties.title = "标擎 BidPilot 零基础操作说明书"
    doc.core_properties.subject = "安装、启动、查询、订阅、配置、停服、备份与排错"
    doc.core_properties.author = "标擎 BidPilot 项目组"
    doc.core_properties.keywords = "BidPilot, 招投标情报, 操作说明书, 长期订阅, LLM"
    doc.core_properties.comments = "由真实项目功能与浏览器验收结果生成"

    add_cover(doc)
    add_front_matter(doc)
    add_body_from_markdown(
        doc,
        source.read_text(encoding="utf-8"),
        assets_dir=assets_dir,
        bullet_abstract_id=bullet_abstract_id,
        number_abstract_id=number_abstract_id,
    )
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description="生成标擎 BidPilot 零基础 Word 说明书")
    parser.add_argument("--input", type=Path, default=Path("docs/BEGINNER_MANUAL.md"))
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("outputs/manuals/标擎BidPilot零基础操作说明书_v0.7.0.docx"),
    )
    parser.add_argument("--assets-dir", type=Path, default=Path("outputs/manuals/assets"))
    args = parser.parse_args()
    build_manual(args.input.resolve(), args.output.resolve(), args.assets_dir.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
