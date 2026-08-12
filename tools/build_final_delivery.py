"""Build the competition-final DOCX set from the official template and repository docs."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

INK = "172033"
BODY = "30394A"
MUTED = "657085"
BLUE = "2F6BFF"
BLUE_DARK = "1F3864"
BLUE_LIGHT = "EAF0FF"
PURPLE = "7B47FF"
PURPLE_LIGHT = "F7F3FF"
ORANGE = "F5A24B"
ORANGE_LIGHT = "FFF5EA"
LINE = "D9DFEA"
WHITE = "FFFFFF"


def set_font(run, size: float, *, bold: bool = False, color: str = BODY) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def ensure_style(
    document: Document,
    name: str,
    *,
    size: float,
    bold: bool = False,
    color: str = BODY,
    before: float = 0,
    after: float = 6,
    line_spacing: float = 1.22,
) -> None:
    try:
        style = document.styles[name]
    except KeyError:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Arial"
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line_spacing


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(
    cell, *, top: int = 80, start: int = 100, bottom: int = 80, end: int = 100
) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = LINE, size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    color: str = BODY,
    size: float = 9.5,
    align=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_font(run, size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    *,
    header_fill: str = BLUE_DARK,
    header_text: str = WHITE,
    first_column_fill: str | None = None,
    widths_cm: list[float] | None = None,
) -> object:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        set_cell_text(table.cell(0, index), header, bold=True, color=header_text, size=9.5)
        shade_cell(table.cell(0, index), header_fill)
        if widths_cm:
            table.cell(0, index).width = Cm(widths_cm[index])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        keep_row_together(table.rows[-1])
        normalized_values = [*values[: len(headers)], *([""] * max(0, len(headers) - len(values)))]
        for column, value in enumerate(normalized_values):
            set_cell_text(cells[column], value, bold=column == 0 and first_column_fill is not None)
            if widths_cm:
                cells[column].width = Cm(widths_cm[column])
            if column == 0 and first_column_fill:
                shade_cell(cells[column], first_column_fill)
            elif row_index % 2:
                shade_cell(cells[column], "F8FAFD")
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    fonts.set(qn("w:eastAsia"), "微软雅黑")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([fonts, color, underline])
    run.append(run_properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph) -> None:
    paragraph.add_run("—  ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    paragraph.add_run("  —")


def configure_document(
    document: Document, *, header_text: str, source_derived: bool = False
) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18 if source_derived else 20)
    section.bottom_margin = Mm(18 if source_derived else 20)
    section.left_margin = Mm(20 if source_derived else 22)
    section.right_margin = Mm(20 if source_derived else 22)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(9)

    for name, args in {
        "Normal": dict(size=10.5, color=BODY, after=6, line_spacing=1.22),
        "Doc Title": dict(size=28, bold=True, color=INK, after=8, line_spacing=1.05),
        "Doc Subtitle": dict(size=13, color=BLUE, after=16, line_spacing=1.1),
        "Heading 1": dict(size=18, bold=True, color=INK, before=12, after=8, line_spacing=1.08),
        "Heading 2": dict(size=14, bold=True, color=BLUE, before=12, after=6, line_spacing=1.1),
        "Heading 3": dict(size=11.5, bold=True, color=INK, before=9, after=4, line_spacing=1.15),
        "Caption": dict(size=8.5, color=MUTED, after=8, line_spacing=1.0),
        "Code": dict(size=8.5, color=INK, after=4, line_spacing=1.0),
    }.items():
        ensure_style(document, name, **args)

    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(header_text)
    set_font(run, 8, bold=True, color=MUTED)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), LINE)
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(paragraph)
    for run in paragraph.runs:
        set_font(run, 8, color=MUTED)

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    # PAGE fields are refreshed by Word's print/export layout.  Forcing every
    # field to update immediately on open can make desktop Word repaginate
    # indefinitely for documents that contain long tables.
    update_fields.set(qn("w:val"), "false")


def add_title(document: Document, text: str, subtitle: str | None = None) -> None:
    paragraph = document.add_paragraph(style="Doc Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(text)
    if subtitle:
        paragraph = document.add_paragraph(style="Doc Subtitle")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(subtitle)


def add_heading(document: Document, text: str, level: int = 1, *, page_break: bool = False) -> None:
    if page_break:
        document.add_page_break()
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    if level == 1:
        marker = paragraph._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "5")
        bottom.set(qn("w:color"), BLUE)
        border.append(bottom)
        marker.append(border)
    set_font(run, {1: 18, 2: 14, 3: 11.5}[level], bold=True, color=INK if level != 2 else BLUE)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="Normal")
        paragraph.paragraph_format.left_indent = Cm(0.7)
        paragraph.paragraph_format.first_line_indent = Cm(-0.45)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(f"•  {item}")
        set_font(run, 10.3)


def add_callout(
    document: Document,
    title: str,
    body: str,
    *,
    fill: str = PURPLE_LIGHT,
    accent: str = PURPLE,
) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color=accent, size="6")
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, top=120, start=150, bottom=120, end=150)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(title)
    set_font(run, 11, bold=True, color=accent)
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.18
    run = paragraph.add_run(body)
    set_font(run, 10)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def strip_inline_markdown(text: str) -> str:
    text = re.sub(r"\[([^]]+)]\(([^)]+)\)", r"\1（\2）", text)
    text = text.replace("`", "").replace("**", "").replace("__", "")
    return text.strip()


def parse_table(lines: list[str]) -> tuple[list[str], list[list[str]]]:
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    if len(rows) >= 2 and all(
        re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in rows[1]
    ):
        rows.pop(1)
    return rows[0], rows[1:]


def render_mermaid_png(code: str, output_dir: Path, index: int) -> Path | None:
    """Render a Mermaid block when the local Mermaid CLI is available.

    Delivery documents should show the four important flows as diagrams, not
    expose raw Mermaid source. If the optional CLI is unavailable, the caller
    keeps a readable text fallback instead of failing the whole release.
    """
    executable = shutil.which("mmdc.cmd") or shutil.which("mmdc")
    if not executable:
        return None
    output_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="bidpilot-mermaid-") as temp_dir:
        source_path = Path(temp_dir) / f"diagram-{index}.mmd"
        target_path = output_dir / f"diagram-{index}.png"
        source_path.write_text(code.strip() + "\n", encoding="utf-8")
        command = [
            executable,
            "-i",
            str(source_path),
            "-o",
            str(target_path),
            "-b",
            "white",
            "--quiet",
        ]
        try:
            subprocess.run(command, check=True, capture_output=True, text=True, timeout=90)
        except (OSError, subprocess.SubprocessError):
            target_path.unlink(missing_ok=True)
            return None
    return target_path if target_path.exists() else None


def add_markdown_image(document: Document, source: Path, alt: str) -> None:
    if source.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        picture = run.add_picture(str(source), width=Cm(16.2))
        accessible_name = alt or source.stem
        picture._inline.docPr.set("descr", accessible_name)
        picture._inline.docPr.set("title", accessible_name)
        caption = document.add_paragraph(style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run(accessible_name)
    else:
        add_callout(document, "素材缺失", f"未找到图片：{source}")


def add_markdown_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pattern = re.compile(r"\[([^]]+)]\((https?://[^)]+)\)|(https?://[^\s]+)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(
                text[cursor : match.start()].replace("`", "").replace("**", "").replace("__", "")
            )
            set_font(run, 10.5)
        label = match.group(1) or match.group(3)
        url = match.group(2) or match.group(3)
        add_hyperlink(paragraph, label, url)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:].replace("`", "").replace("**", "").replace("__", ""))
        set_font(run, 10.5)


def build_markdown_docx(
    source: Path,
    output: Path,
    *,
    subtitle: str,
    asset_root: Path | None = None,
    template: Path | None = None,
    diagram_cache_dir: Path | None = None,
) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    template_section = None
    if template:
        if not template.exists():
            raise FileNotFoundError(template)
        # Read the retained template as a geometry reference only.  Reusing an
        # already-filled DOCX package can carry stale controls, relationships,
        # drawings and field caches into the new file and make desktop Word
        # repaginate indefinitely.  A clean package keeps the official page
        # system without inheriting those opaque objects.
        template_document = Document(str(template))
        template_section = template_document.sections[0]
    document = Document()
    configure_document(
        document,
        header_text=f"聚标成擎｜{subtitle}",
        source_derived=template_section is not None,
    )
    if template_section is not None:
        section = document.sections[0]
        section.page_width = template_section.page_width
        section.page_height = template_section.page_height
        section.orientation = template_section.orientation
        section.top_margin = template_section.top_margin
        section.bottom_margin = template_section.bottom_margin
        section.left_margin = template_section.left_margin
        section.right_margin = template_section.right_margin
        section.header_distance = template_section.header_distance
        section.footer_distance = template_section.footer_distance
    document.core_properties.author = "聚标成擎"
    title = next((line[2:].strip() for line in lines if line.startswith("# ")), source.stem)
    add_title(document, title, subtitle)
    index = 0
    in_code = False
    code_lines: list[str] = []
    code_language = ""
    diagram_index = 0
    first_title_consumed = False
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            if in_code:
                if code_language in {"mermaid", "mermaid.js"}:
                    diagram_index += 1
                    if diagram_cache_dir:
                        cached_image = diagram_cache_dir / f"diagram-{diagram_index}.png"
                        if not cached_image.exists():
                            raise FileNotFoundError(cached_image)
                        manifest_path = diagram_cache_dir.parent / "manifest.json"
                        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
                        expected_hash = manifest.get(diagram_cache_dir.name, {}).get(
                            cached_image.name
                        )
                        actual_hash = hashlib.sha256(
                            "\n".join(code_lines).strip().encode("utf-8")
                        ).hexdigest()
                        if expected_hash != actual_hash:
                            raise RuntimeError(
                                f"Mermaid 图缓存与源稿不一致：{cached_image}; "
                                "请重新渲染并更新 manifest.json"
                            )
                        image = cached_image
                    else:
                        image = render_mermaid_png(
                            "\n".join(code_lines),
                            output.parent / ".diagram-assets",
                            diagram_index,
                        )
                    if image:
                        add_markdown_image(document, image, f"图 {diagram_index}｜流程图")
                    else:
                        table = document.add_table(rows=1, cols=1)
                        set_table_borders(table, color=LINE)
                        cell = table.cell(0, 0)
                        shade_cell(cell, "F4F6F8")
                        set_cell_text(cell, "\n".join(code_lines), size=8.2, color=INK)
                else:
                    table = document.add_table(rows=1, cols=1)
                    set_table_borders(table, color=LINE)
                    cell = table.cell(0, 0)
                    shade_cell(cell, "F4F6F8")
                    set_cell_text(cell, "\n".join(code_lines), size=8.2, color=INK)
                code_lines = []
                code_language = ""
                in_code = False
            else:
                in_code = True
                code_language = line[3:].strip().lower()
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if line.startswith("# "):
            if not first_title_consumed:
                first_title_consumed = True
            else:
                add_heading(document, strip_inline_markdown(line[2:]), page_break=True)
        elif line.startswith("## "):
            add_heading(document, strip_inline_markdown(line[3:]), level=1)
        elif line.startswith("### "):
            add_heading(document, strip_inline_markdown(line[4:]), level=2)
        elif line.startswith("#### "):
            add_heading(document, strip_inline_markdown(line[5:]), level=3)
        elif line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            headers, rows = parse_table(table_lines)
            add_table(
                document,
                [strip_inline_markdown(item) for item in headers],
                [[strip_inline_markdown(item) for item in row] for row in rows],
            )
            continue
        elif re.match(r"^\s*[-*] ", line):
            add_bullets(document, [strip_inline_markdown(re.sub(r"^\s*[-*] ", "", line))])
        elif re.match(r"^\s*\d+[.)] ", line):
            number = re.match(r"^\s*(\d+)[.)] ", line).group(1)
            paragraph = document.add_paragraph(style="Normal")
            paragraph.paragraph_format.left_indent = Cm(0.7)
            paragraph.paragraph_format.first_line_indent = Cm(-0.45)
            item = re.sub(r"^\s*\d+[.)] ", "", line)
            run = paragraph.add_run(f"{number}.  {strip_inline_markdown(item)}")
            set_font(run, 10.3)
        elif line.startswith("> "):
            add_callout(document, "说明", strip_inline_markdown(line[2:]))
        elif image_match := re.match(r"^!\[([^]]*)\]\(([^)]+)\)$", line.strip()):
            raw_path = image_match.group(2).strip()
            image_path = (source.parent / raw_path).resolve()
            if not image_path.exists() and asset_root:
                image_path = (asset_root / raw_path).resolve()
            add_markdown_image(document, image_path, image_match.group(1).strip())
        elif line.strip() and not line.strip() == "---":
            add_markdown_paragraph(document, line.strip())
        index += 1
    # Long tables can make desktop Word spend unbounded time repaginating when
    # every preceding heading is chained to the next block.  The generated
    # documents use explicit hierarchy and page breaks, so relax that chain.
    for paragraph in document.paragraphs:
        if paragraph.style and paragraph.style.name.startswith("Heading "):
            paragraph.paragraph_format.keep_with_next = False
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def copy_reports(source_dir: Path, target_dir: Path) -> None:
    target_dir.mkdir(parents=True, exist_ok=True)
    pattern = re.compile(r"_(\d{12})_[0-9a-f]{8}(?:-[^.]*)?\.docx$", re.IGNORECASE)
    for source in sorted(source_dir.glob("*.docx")):
        target_name = pattern.sub(r"_\1.docx", source.name)
        shutil.copy2(source, target_dir / target_name)


def copy_incremental_evidence(source_dir: Path, reports_dir: Path, audit_dir: Path) -> int:
    """Copy a complete, isolated two-run subscription audit into the delivery package."""

    summary_path = source_dir / "00_incremental_evidence_summary.json"
    if not summary_path.exists():
        raise FileNotFoundError(summary_path)
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    assertions = summary.get("assertions", {})
    required_assertions = {
        "database_integrity_ok",
        "first_report_exists",
        "ledger_did_not_grow_on_second_run",
        "same_subscription_id_for_both_runs",
        "second_outbox_is_auditable_skip",
        "second_run_no_new_report",
        "second_run_zero_new",
        "two_persisted_runs",
    }
    failed = sorted(name for name in required_assertions if assertions.get(name) is not True)
    if failed:
        raise RuntimeError("真实增量证据断言未通过：" + "、".join(failed))

    report_source = source_dir / "reports"
    reports = sorted(report_source.glob("*.docx"))
    if len(reports) != 1:
        raise RuntimeError("增量证据应恰好包含首轮 Word；第二轮零新增不应生成空报告")
    copy_reports(report_source, reports_dir)

    evidence_target = audit_dir / "真实订阅增量证据"
    evidence_target.mkdir(parents=True, exist_ok=True)
    evidence_files = (
        "00_incremental_evidence_summary.json",
        "01_intent_parse.json",
        "02_subscription_created.json",
        "03_first_run_response.json",
        "04_after_first_database_snapshot.json",
        "05_second_run_response.json",
        "06_after_second_database_snapshot.json",
        "07_api_audit_endpoints.json",
        "08_file_manifest.json",
        "incremental_evidence.db",
        "run_incremental_evidence.py",
        "SHA256SUMS.txt",
        "证据说明.txt",
    )
    for name in evidence_files:
        source = source_dir / name
        if not source.exists():
            raise FileNotFoundError(source)
        shutil.copy2(source, evidence_target / name)
    return len(reports)


def copy_operation_documents(repo: Path, target_dir: Path) -> None:
    target_dir.mkdir(parents=True, exist_ok=True)
    mapping = {
        repo / "README.md": "README.md",
        repo / "docs" / "API_REFERENCE.md": "API参考.md",
        repo / "docs" / "CONFIGURATION_GUIDE.md": "配置指南.md",
        repo / "docs" / "USER_GUIDE.md": "用户指南.md",
        repo / "docs" / "BEGINNER_MANUAL.md": "零基础操作说明.md",
        repo / "docs" / "比赛验收指南.md": "比赛验收指南.md",
        repo / "docs" / "交接手册_v0.8.0.md": "交接手册_v0.8.0.md",
        repo / "docs" / "DELIVERY_CHECKLIST.md": "发布验收清单.md",
    }
    for source, name in mapping.items():
        if not source.exists():
            raise FileNotFoundError(source)
        shutil.copy2(source, target_dir / name)


def copy_beginner_manual(repo: Path, target_dir: Path) -> None:
    target_dir.mkdir(parents=True, exist_ok=True)
    manuals = repo / "outputs" / "manuals"
    for suffix in ("docx", "pdf"):
        source = manuals / f"标擎BidPilot零基础操作说明书_v0.8.0.{suffix}"
        if not source.exists():
            raise FileNotFoundError(source)
        shutil.copy2(source, target_dir / source.name)


def copy_ui_evidence(repo: Path, target_dir: Path) -> None:
    target_dir.mkdir(parents=True, exist_ok=True)
    assets = repo / "outputs" / "manuals" / "assets"
    for name in ("12-home-1024.png", "13-home-390.png", "14-results-1440.png"):
        source = assets / name
        if not source.exists():
            raise FileNotFoundError(source)
        shutil.copy2(source, target_dir / name)


def build_source_archive(repo: Path, target: Path) -> None:
    """Archive exactly the committed tree; ignored credentials and generated data stay out."""

    target.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["git", "archive", "--format=zip", f"--output={target}", "HEAD"],
        cwd=repo,
        check=True,
    )


def write_delivery_manifests(output: Path) -> None:
    audit_dir = output / "05_验收与校验"
    audit_dir.mkdir(parents=True, exist_ok=True)
    file_list = audit_dir / "文件清单.txt"
    checksum_list = audit_dir / "SHA256SUMS.txt"

    relative_files = sorted(
        path.relative_to(output).as_posix()
        for path in output.rglob("*")
        if path.is_file() and path not in {file_list, checksum_list}
    )
    relative_files.extend(
        [
            file_list.relative_to(output).as_posix(),
            checksum_list.relative_to(output).as_posix(),
        ]
    )
    file_list.write_text("\n".join(sorted(relative_files)) + "\n", encoding="utf-8")

    checksum_lines: list[str] = []
    for path in sorted(item for item in output.rglob("*") if item.is_file()):
        if path == checksum_list:
            continue
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        checksum_lines.append(f"{digest}  {path.relative_to(output).as_posix()}")
    checksum_list.write_text("\n".join(checksum_lines) + "\n", encoding="utf-8")


def git_text(repo: Path, *args: str) -> str:
    return subprocess.run(
        ["git", *args], cwd=repo, check=True, capture_output=True, text=True, encoding="utf-8"
    ).stdout.strip()


def write_build_metadata(
    repo: Path, output: Path, report_count: int, *, incremental_evidence_included: bool
) -> None:
    audit_dir = output / "05_验收与校验"
    audit_dir.mkdir(parents=True, exist_ok=True)
    status = git_text(repo, "status", "--short")
    metadata = {
        "built_at": datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(timespec="seconds"),
        "git_commit": git_text(repo, "rev-parse", "HEAD"),
        "git_branch": git_text(repo, "branch", "--show-current"),
        "git_worktree_clean": not bool(status),
        "report_count": report_count,
        "incremental_evidence_included": incremental_evidence_included,
        "source_archive_scope": "git archive HEAD",
        "notes": [
            "测试数量与 CI 结论以同目录跨平台测试与审计报告为准。",
            "公开视频链接只有在真实上传并验证后才会标记完成。",
        ],
    }
    (audit_dir / "构建信息.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def prepare_video_delivery(
    demo_dir: Path,
    video_link_file: Path | None,
    demo_video: Path | None = None,
) -> tuple[bool, bool]:
    demo_dir.mkdir(parents=True, exist_ok=True)
    local_video_ready = False
    if demo_video is not None:
        source = demo_video.resolve()
        if not source.is_file():
            raise FileNotFoundError(f"Demo 视频不存在：{source}")
        if source.suffix.lower() != ".mp4" or source.stat().st_size < 1024:
            raise ValueError(f"Demo 视频必须是非空 MP4 文件：{source}")
        with source.open("rb") as handle:
            header = handle.read(32)
        if b"ftyp" not in header:
            raise ValueError(f"Demo 视频缺少 MP4 文件头：{source}")
        shutil.copy2(source, demo_dir / source.name)
        local_video_ready = True

    target = demo_dir / "视频公开链接.txt"
    if video_link_file and video_link_file.exists():
        content = video_link_file.read_text(encoding="utf-8").strip()
        if re.search(r"https?://", content) and not re.search(
            r"待|占位|TODO|example", content, re.I
        ):
            target.write_text(content + "\n", encoding="utf-8")
            return local_video_ready, True
    pending_message = (
        "本地 Demo 视频已完成并随交付包提供；待团队上传后，使用未登录窗口验证公开可访问，再把真实链接回填到本文件。\n"
        if local_video_ready
        else "待团队完成真实录制、上传并使用未登录窗口验证后回填。\n"
    )
    target.write_text(
        pending_message + "本文件不提供虚构公开视频链接。\n",
        encoding="utf-8",
    )
    return local_video_ready, False


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--repo", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--reports-source", type=Path)
    parser.add_argument("--incremental-evidence", type=Path)
    parser.add_argument("--previous-delivery", type=Path)
    parser.add_argument("--video-link-file", type=Path)
    parser.add_argument("--demo-video", type=Path)
    args = parser.parse_args()

    repo = args.repo.resolve()
    output = args.output.resolve()
    if output.exists() and any(output.iterdir()):
        raise RuntimeError(f"输出目录必须为空，避免旧文件混入新交付包：{output}")
    output.mkdir(parents=True, exist_ok=True)
    competition = (
        output
        / "06_参赛方案"
        / "【40强赛】超聚变🤝聚标成擎｜标擎BidPilot——一句话查标讯，原文能复核，订阅只发新内容.docx"
    )
    # The proposal is authored as readable Markdown so product language and
    # evidence can be reviewed in Git.  The supplied 40-strong template
    # provides the page geometry; content and styles are rebuilt in a clean
    # OOXML package so stale placeholders or controls cannot survive.
    build_markdown_docx(
        repo / "docs" / "参赛方案.md",
        competition,
        subtitle="超聚变命题｜参赛方案",
        asset_root=repo,
        template=args.template.resolve(),
        diagram_cache_dir=repo / "outputs" / "manuals" / "assets" / "diagrams" / "proposal",
    )
    shutil.copy2(repo / "docs" / "参赛方案.md", competition.parent / "参赛方案_可编辑源稿.md")

    design_dir = output / "02_设计文档"
    build_markdown_docx(
        repo / "docs" / "详设文档.md",
        design_dir / "详设文档.docx",
        subtitle="详细设计说明书",
        asset_root=repo,
        diagram_cache_dir=repo / "outputs" / "manuals" / "assets" / "diagrams" / "design",
    )
    build_markdown_docx(
        repo / "docs" / "LINUX_DEPLOYMENT.md",
        design_dir / "Linux部署与运维手册.docx",
        subtitle="Linux 生产部署、登录来源与运维",
    )
    audit_dir = output / "05_验收与校验"
    build_markdown_docx(
        repo / "docs" / "CROSS_PLATFORM_TEST_AUDIT.md",
        audit_dir / "跨平台测试与审计报告.docx",
        subtitle="Windows / Linux / macOS / Container",
    )
    copy_ui_evidence(repo, audit_dir / "UI响应式验收截图")
    checklist_dir = output / "00_提交说明与清单"
    build_markdown_docx(
        repo / "docs" / "DELIVERY_CHECKLIST.md",
        checklist_dir / "提交说明与验收清单.docx",
        subtitle="超聚变命题终版交付",
    )

    reports_source = args.reports_source
    if reports_source is None and args.previous_delivery:
        reports_source = args.previous_delivery / "01_运行结果Word"
    if reports_source is None or not reports_source.exists():
        raise FileNotFoundError("请使用 --reports-source 指定当前版本真实运行结果目录")
    copy_reports(reports_source.resolve(), output / "01_运行结果Word")
    incremental_evidence_included = False
    if args.incremental_evidence:
        copy_incremental_evidence(
            args.incremental_evidence.resolve(),
            output / "01_运行结果Word",
            audit_dir,
        )
        incremental_evidence_included = True
    report_count = len(list((output / "01_运行结果Word").glob("*.docx")))
    if report_count < 2:
        raise RuntimeError("官方要求多个问题的运行结果 Word；当前复制数量少于 2")

    for source in (
        repo / "docs" / "详设文档.md",
        repo / "docs" / "FINALIST_BENCHMARK.md",
        repo / "SPEC.md",
    ):
        shutil.copy2(source, design_dir / source.name)
    for source in (
        repo / "docs" / "LINUX_DEPLOYMENT.md",
        repo / "docs" / "CROSS_PLATFORM_TEST_AUDIT.md",
    ):
        target = design_dir if "LINUX" in source.name else audit_dir
        shutil.copy2(source, target / source.name)

    code_dir = output / "03_代码与操作文档"
    copy_beginner_manual(repo, code_dir)
    copy_operation_documents(repo, code_dir / "操作指南")
    build_source_archive(repo, code_dir / "标擎BidPilot_v0.8.0_源码.zip")

    demo_dir = output / "04_Demo"
    demo_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(repo / "docs" / "COMPETITION_DEMO_SCRIPT.md", demo_dir / "Demo录制脚本与分镜.md")
    local_video_ready, public_link_ready = prepare_video_delivery(
        demo_dir,
        args.video_link_file,
        args.demo_video,
    )

    for assets in output.rglob(".diagram-assets"):
        if assets.is_dir():
            shutil.rmtree(assets)

    write_build_metadata(
        repo,
        output,
        report_count,
        incremental_evidence_included=incremental_evidence_included,
    )
    if local_video_ready and public_link_ready:
        status_text = "四类官方交付物已齐备；本地 Demo 视频已纳入，公开视频链接已验证。"
    elif local_video_ready:
        status_text = "本地 Demo 视频已纳入交付包；公开视频仍需团队真实上传并回填链接。"
    elif public_link_ready:
        status_text = "四类官方交付物已齐备；公开视频链接已验证。"
    else:
        status_text = "代码、设计文档和多份运行结果已生成；Demo 视频仍需团队真实录制并回填链接。"
    (output / "00_提交说明与清单" / "当前交付状态.txt").write_text(
        status_text + "\n", encoding="utf-8"
    )

    write_delivery_manifests(output)

    print(competition)


if __name__ == "__main__":
    main()
